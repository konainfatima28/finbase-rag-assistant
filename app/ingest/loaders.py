"""Document loaders (PROMPT.md §5.1): registry by file extension, all emitting `RawDocument`.

PDF: PyMuPDF for text lines + page numbers + geometry, pdfplumber for tables. Text lines that fall inside
a detected table's bounding box are replaced by a single `⟦TABLE:id⟧` placeholder line at the table's
reading position, so tables are never duplicated as loose cell text. The rupee sign in these PDFs is a
ZapfDingbats glyph; we map every ZapfDingbats char to '■' (the canonical corrupted form) and let the
cleaner's single rule convert '■'+digit -> '₹'.
"""

from __future__ import annotations

import csv
import hashlib
import io
from collections.abc import Callable
from pathlib import Path
from typing import Any

import structlog

from app.ingest.cleaner import clean_cell
from app.ingest.models import DocMeta, Line, Page, RawDocument, Table

log = structlog.get_logger(__name__)

GLYPH = "■"
_WRAP_TOLERANCE_PT = 40.0
_TOP_OF_PAGE_PT = 60.0
_BBOX_PAD_PT = 1.5

Loader = Callable[[Path, DocMeta], RawDocument]
_REGISTRY: dict[str, Loader] = {}


def register(*extensions: str) -> Callable[[Loader], Loader]:
    """Register a loader for file extensions (lower-case, with dot)."""

    def deco(fn: Loader) -> Loader:
        for ext in extensions:
            _REGISTRY[ext] = fn
        return fn

    return deco


def load_document(path: Path, meta: DocMeta) -> RawDocument:
    """Load any supported file into a `RawDocument`."""
    loader = _REGISTRY.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"unsupported file type: {path.suffix} ({path.name})")
    return loader(path, meta)


def supported_extensions() -> list[str]:
    """Extensions with a registered loader."""
    return sorted(_REGISTRY)


def file_sha256(path: Path) -> str:
    """Hex sha256 of a file."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def table_placeholder(table_id: str) -> str:
    """Placeholder text standing for a table in the line stream."""
    return f"⟦TABLE:{table_id}⟧"


# --- PDF -------------------------------------------------------------------------------------------
def _is_zapf(fontname: str) -> bool:
    return "zapfdingbats" in fontname.lower()


PdfTable = tuple[tuple[float, float, float, float], list[list[str]], list[str]]


def _pdf_tables(page: Any) -> list[PdfTable]:
    for char in page.chars:
        if _is_zapf(str(char.get("fontname", ""))):
            char["text"] = GLYPH
    found: list[PdfTable] = []
    for table in page.find_tables():
        raw_rows = [[c or "" for c in row] for row in table.extract()]
        merged = [clean_cell(c) for row in raw_rows for c in row if "\n" in c.strip()]
        rows = [[clean_cell(c) for c in row] for row in raw_rows]
        rows = [r for r in rows if any(r)]
        if rows:
            found.append((tuple(table.bbox), rows, merged))
    return found


def _mostly_bold(spans: list[dict[str, Any]]) -> bool:
    """True if most visible (non-bullet, non-glyph) characters of a line are set in a bold font."""
    bold = regular = 0
    for span in spans:
        if _is_zapf(span["font"]):
            continue
        visible = sum(1 for c in span["chars"] if c["c"].strip() and c["c"] != "•")
        if "bold" in span["font"].lower():
            bold += visible
        else:
            regular += visible
    return bold > regular


def _inside(bbox: tuple[float, float, float, float], box: tuple[float, float, float, float]) -> bool:
    cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
    return (
        box[0] - _BBOX_PAD_PT <= cx <= box[2] + _BBOX_PAD_PT
        and box[1] - _BBOX_PAD_PT <= cy <= box[3] + _BBOX_PAD_PT
    )


@register(".pdf")
def load_pdf(path: Path, meta: DocMeta) -> RawDocument:
    """Load a text-layer PDF (OCR fallback only for pages without any text)."""
    import pdfplumber
    import pymupdf

    pages: list[Page] = []
    lines: list[Line] = []
    tables: dict[str, Table] = {}
    last_table: Table | None = None  # for cross-page continuation
    last_element_was_table = False
    with pymupdf.open(path) as mdoc, pdfplumber.open(path) as pdoc:  # type: ignore[no-untyped-call]
        for index, (mpage, ppage) in enumerate(zip(mdoc, pdoc.pages, strict=True)):
            page_no = index + 1
            ptables = _pdf_tables(ppage)
            raw_lines: list[tuple[tuple[float, float, float, float], str]] = []
            bold_by_bbox: dict[tuple[float, float, float, float], bool] = {}
            for block in mpage.get_text("rawdict")["blocks"]:
                for mline in block.get("lines", []):
                    text = "".join(
                        GLYPH * len(span["chars"])
                        if _is_zapf(span["font"])
                        else "".join(c["c"] for c in span["chars"])
                        for span in mline["spans"]
                    )
                    if text.strip():
                        bbox = tuple(mline["bbox"])
                        bold_by_bbox[bbox] = _mostly_bold(mline["spans"])
                        raw_lines.append((bbox, text))
            if not raw_lines and not ptables:
                raw_lines = _ocr_page(mpage, page_no, path)
            pages.append(Page(number=page_no, raw_text="\n".join(t for _, t in raw_lines)))
            right_edge = max((b[2] for b, _ in raw_lines), default=0.0)

            elements: list[tuple[float, float, Line | PdfTable]] = []
            for bbox, text in raw_lines:
                if any(_inside(bbox, tb) for tb, _, _ in ptables):
                    continue
                wrapped = bbox[2] >= right_edge - _WRAP_TOLERANCE_PT
                line = Line(text=text, page=page_no, wrapped=wrapped, bold=bold_by_bbox.get(bbox))
                elements.append((bbox[1], bbox[0], line))
            for ptable in ptables:
                elements.append((ptable[0][1], ptable[0][0], ptable))
            elements.sort(key=lambda e: (round(e[0], 1), e[1]))

            for pos, (_y, _x, element) in enumerate(elements):
                if isinstance(element, Line):
                    lines.append(element)
                    last_element_was_table = False
                    continue
                tb, rows, merged = element
                is_continuation = (
                    pos == 0
                    and tb[1] < _TOP_OF_PAGE_PT
                    and last_element_was_table
                    and last_table is not None
                    and len(rows[0]) == len(last_table.header)
                )
                if is_continuation and last_table is not None:
                    body = rows[1:] if rows[0] == last_table.header else rows
                    last_table.rows.extend(body)
                    last_table.merged_cells.extend(merged)
                    last_table.page_end = page_no
                    last_element_was_table = True
                    continue
                table_id = f"{meta.doc_id}-t{len(tables) + 1:03d}"
                table = Table(
                    table_id=table_id,
                    page_start=page_no,
                    page_end=page_no,
                    header=rows[0],
                    rows=rows[1:],
                    merged_cells=merged,
                )
                tables[table_id] = table
                lines.append(Line(text=table_placeholder(table_id), page=page_no, table_id=table_id))
                last_table = table
                last_element_was_table = True
    return RawDocument(
        meta=meta, source_path=path, sha256=file_sha256(path), pages=pages, lines=lines, tables=tables
    )


def _ocr_page(mpage: Any, page_no: int, path: Path) -> list[tuple[tuple[float, float, float, float], str]]:
    """OCR fallback for image-only pages (requires Tesseract; logs and returns [] if unavailable)."""
    try:
        textpage = mpage.get_textpage_ocr(full=True)
        text = mpage.get_text(textpage=textpage)
    except Exception as exc:  # Tesseract missing or OCR failure: degrade, never crash ingestion
        log.warning("ocr_unavailable", file=path.name, page=page_no, error=str(exc))
        return []
    width = float(mpage.rect.width)
    return [((0.0, float(i), width, float(i) + 1.0), t) for i, t in enumerate(text.splitlines()) if t.strip()]


# --- plain text / markdown ------------------------------------------------------------------------
@register(".txt", ".md")
def load_text(path: Path, meta: DocMeta) -> RawDocument:
    """Load .txt/.md; markdown heading hashes are dropped (text kept), pages = form-feed separated."""
    content = path.read_text(encoding="utf-8")
    pages: list[Page] = []
    lines: list[Line] = []
    for index, page_text in enumerate(content.split("\f")):
        page_no = index + 1
        pages.append(Page(number=page_no, raw_text=page_text))
        for raw in page_text.splitlines():
            text = raw.lstrip("#").strip() if raw.lstrip().startswith("#") else raw.strip()
            if text:
                lines.append(Line(text=text, page=page_no))
    return RawDocument(meta=meta, source_path=path, sha256=file_sha256(path), pages=pages, lines=lines)


# --- CSV -----------------------------------------------------------------------------------------
@register(".csv")
def load_csv(path: Path, meta: DocMeta) -> RawDocument:
    """Load a CSV as one table (header = first row); chunked as row groups downstream."""
    content = path.read_text(encoding="utf-8-sig")
    rows = [
        [clean_cell(c) for c in row]
        for row in csv.reader(io.StringIO(content))
        if any(c.strip() for c in row)
    ]
    if not rows:
        raise ValueError(f"empty CSV: {path.name}")
    table_id = f"{meta.doc_id}-t001"
    table = Table(table_id=table_id, page_start=1, page_end=1, header=rows[0], rows=rows[1:])
    lines = [Line(text=table_placeholder(table_id), page=1, table_id=table_id)]
    return RawDocument(
        meta=meta,
        source_path=path,
        sha256=file_sha256(path),
        pages=[Page(1, content)],
        lines=lines,
        tables={table_id: table},
    )


# --- DOCX ----------------------------------------------------------------------------------------
@register(".docx")
def load_docx(path: Path, meta: DocMeta) -> RawDocument:
    """Load a .docx: paragraphs and tables in body order (single logical page)."""
    import docx

    document = docx.Document(str(path))
    lines: list[Line] = []
    tables: dict[str, Table] = {}
    raw_parts: list[str] = []
    paragraphs = iter(document.paragraphs)
    docx_tables = iter(document.tables)
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = next(paragraphs).text.strip()
            if text:
                lines.append(Line(text=text, page=1))
                raw_parts.append(text)
        elif tag == "tbl":
            dtable = next(docx_tables)
            rows = [[clean_cell(cell.text) for cell in row.cells] for row in dtable.rows]
            rows = [r for r in rows if any(r)]
            if not rows:
                continue
            table_id = f"{meta.doc_id}-t{len(tables) + 1:03d}"
            tables[table_id] = Table(
                table_id=table_id, page_start=1, page_end=1, header=rows[0], rows=rows[1:]
            )
            lines.append(Line(text=table_placeholder(table_id), page=1, table_id=table_id))
            raw_parts.extend(" | ".join(r) for r in rows)
    return RawDocument(
        meta=meta,
        source_path=path,
        sha256=file_sha256(path),
        pages=[Page(1, "\n".join(raw_parts))],
        lines=lines,
        tables=tables,
    )
